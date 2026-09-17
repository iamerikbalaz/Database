import io
import stat
import zipfile
from xml.sax.saxutils import escape

import pytest

from app import import_workbooks as workbook
from app.import_sources import ImportSourceError

S = workbook.S[1:-1]
R = workbook.R[1:-1]
P = workbook.P[1:-1]
C = workbook.C[1:-1]


def cell(reference, value, kind="inlineStr"):
    content = f"<is><t>{escape(value)}</t></is>" if kind == "inlineStr" else f"<v>{escape(value)}</v>"
    return f'<c r="{reference}" t="{kind}">{content}</c>'


def parts():
    return {
        "[Content_Types].xml": f'<Types xmlns="{C}"><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/></Types>',
        "_rels/.rels": f'<Relationships xmlns="{P}"><Relationship Id="root" Type="{workbook.REL}officeDocument" Target="xl/workbook.xml"/></Relationships>',
        "xl/workbook.xml": f'<workbook xmlns="{S}" xmlns:r="{R}"><sheets><sheet name="Materials" sheetId="1" r:id="first"/></sheets></workbook>',
        "xl/_rels/workbook.xml.rels": f'<Relationships xmlns="{P}"><Relationship Id="first" Type="{workbook.REL}worksheet" Target="worksheets/sheet1.xml"/></Relationships>',
        "xl/worksheets/sheet1.xml": f'<worksheet xmlns="{S}"><sheetData><row r="1">{cell("A1", "Identity")}{cell("B1", "Name")}{cell("C1", "Note")}</row><row r="2">{cell("A2", "RWT_0007_WOOD")}{cell("B2", "Světlý dub")}</row></sheetData></worksheet>',
    }


def archive(values=None, extra=()):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as output:
        for name, value in (parts() if values is None else values).items():
            output.writestr(name, value)
        for name, value in extra:
            output.writestr(name, value)
    return stream.getvalue()


def read(values=None):
    return workbook.read_xlsx(archive(values), sheet="Materials")


def test_reads_literal_values_from_an_explicit_sheet_and_preserves_sparse_columns():
    data = archive()
    assert workbook.xlsx_sheets(data) == ("Materials",)
    table = workbook.read_xlsx(data, sheet="Materials")
    assert table.headers == ("Identity", "Name", "Note")
    assert table.rows[0].number == 2
    assert table.rows[0].values == ("RWT_0007_WOOD", "Světlý dub", "")
    with pytest.raises(ImportSourceError, match="IMPORT_SHEET_REQUIRED"):
        workbook.read_xlsx(data, sheet="PRIVATE_SYNTHETIC_MISSING")


def test_shared_strings_rich_text_and_stored_numeric_values_without_format_guessing():
    values = parts()
    values["xl/sharedStrings.xml"] = f'<sst xmlns="{S}"><si><r><t>Český </t></r><r><t>dub</t></r><rPh sb="0" eb="2"><t>ignored annotation</t></rPh></si></sst>'
    values["xl/worksheets/sheet1.xml"] = values["xl/worksheets/sheet1.xml"].replace(cell("B2", "Světlý dub"), cell("B2", "0", "s") + cell("C2", "12.3400", "n"))
    assert read(values).rows[0].values == ("RWT_0007_WOOD", "Český dub", "12.3400")


@pytest.mark.parametrize(("kind", "value", "expected"), [("b", "1", "TRUE"), ("b", "0", "FALSE"), ("d", "2026-09-17", "2026-09-17"), ("str", "Literal", "Literal")])
def test_other_inert_cell_values(kind, value, expected):
    values = parts()
    values["xl/worksheets/sheet1.xml"] = values["xl/worksheets/sheet1.xml"].replace(cell("B2", "Světlý dub"), cell("B2", value, kind))
    assert read(values).rows[0].values[1] == expected


@pytest.mark.parametrize(("kind", "value"), [("s", "999"), ("s", "-1"), ("e", "#REF!"), ("b", "2"), ("n", "NaN"), ("n", "INF"), ("unknown", "1"), ("n", "PRIVATE_SYNTHETIC")])
def test_invalid_cells_fail_without_disclosing_values(kind, value):
    values = parts()
    values["xl/worksheets/sheet1.xml"] = values["xl/worksheets/sheet1.xml"].replace(cell("B2", "Světlý dub"), cell("B2", value, kind))
    with pytest.raises(ImportSourceError) as caught:
        read(values)
    assert str(caught.value) == "IMPORT_XLSX_CELL"
    assert caught.value.row == 2


def test_never_trusts_formula_caches_in_selected_or_unselected_sheets():
    for selected in (True, False):
        values = parts()
        content = f'<worksheet xmlns="{S}"><sheetData><row r="1"><c r="A1"><f>PRIVATE_SYNTHETIC_FORMULA</f><v>123</v></c></row></sheetData></worksheet>'
        values["xl/worksheets/sheet1.xml" if selected else "xl/worksheets/unused.xml"] = content
        with pytest.raises(ImportSourceError, match="^IMPORT_FORMULA$"):
            read(values)


@pytest.mark.parametrize("encoding", ["utf-8", "utf-16"])
def test_rejects_dtd_entities_before_any_resolution(encoding, monkeypatch):
    import socket
    monkeypatch.setattr(socket, "create_connection", lambda *args, **kwargs: pytest.fail("XML must never open a connection"))
    values = parts()
    values["xl/worksheets/sheet1.xml"] = ('<?xml version="1.0" encoding="' + encoding + '"?><!DOCTYPE x [<!ENTITY secret SYSTEM "file:///PRIVATE_SYNTHETIC_PATH">]><x>&secret;</x>').encode(encoding)
    with pytest.raises(ImportSourceError, match="^IMPORT_XLSX_INVALID$"):
        read(values)


def test_rejects_external_relationships_even_in_unused_parts():
    values = parts()
    values["xl/worksheets/_rels/sheet1.xml.rels"] = f'<Relationships xmlns="{P}"><Relationship Id="external" Type="{workbook.REL}hyperlink" Target="https://example.invalid/PRIVATE_SYNTHETIC" TargetMode="External"/></Relationships>'
    with pytest.raises(ImportSourceError, match="^IMPORT_EXTERNAL_LINK$"):
        read(values)


@pytest.mark.parametrize("name", ["xl/vbaProject.bin", "xl/activeX/control.xml", "xl/embeddings/object.xml", "xl/externalLinks/externalLink1.xml", "xl/connections.xml"])
def test_rejects_macros_embedded_objects_and_connections(name):
    values = parts()
    values[name] = b"PRIVATE_SYNTHETIC"
    with pytest.raises(ImportSourceError, match="^IMPORT_ACTIVE_CONTENT$"):
        read(values)


@pytest.mark.parametrize("name", ["../outside.xml", "/outside.xml", "xl\\outside.xml", "C:/outside.xml", "xl/%2e%2e/outside.xml"])
def test_rejects_unsafe_archive_names_without_extraction(name, monkeypatch):
    monkeypatch.setattr(zipfile.ZipFile, "extract", lambda *args, **kwargs: pytest.fail("No extraction"))
    monkeypatch.setattr(zipfile.ZipFile, "extractall", lambda *args, **kwargs: pytest.fail("No extraction"))
    # Windows' ZIP writer normalizes backslashes. Modify both serialized header
    # names so the reader receives the actual unsafe archive, on every platform.
    encoded_name = name.replace("\\", "/")
    data = archive(extra=[(encoded_name, "<x/>")]).replace(encoded_name.encode(), name.encode())
    with pytest.raises(ImportSourceError, match="^IMPORT_ZIP_ENTRY$"):
        workbook.read_xlsx(data, sheet="Materials")


def test_rejects_duplicate_archive_names_and_symlinks():
    with pytest.warns(UserWarning, match="Duplicate name"):
        duplicate = archive(extra=[("xl/workbook.xml", parts()["xl/workbook.xml"])])
    with pytest.raises(ImportSourceError, match="^IMPORT_ZIP_ENTRY$"):
        workbook.read_xlsx(duplicate, sheet="Materials")
    link = zipfile.ZipInfo("xl/link.xml")
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    with pytest.raises(ImportSourceError, match="^IMPORT_ZIP_ENTRY$"):
        workbook.read_xlsx(archive(extra=[(link, "../private")]), sheet="Materials")


def test_rejects_expansion_and_structure_limits_before_processing_cells(monkeypatch):
    monkeypatch.setattr(workbook, "MAX_PART_BYTES", 2048)
    values = parts()
    values["xl/oversize.xml"] = " " * 2049
    with pytest.raises(ImportSourceError, match="^IMPORT_ZIP_LIMIT$"):
        read(values)
    values = parts()
    values["xl/nested.xml"] = "<x>" * 65 + "</x>" * 65
    with pytest.raises(ImportSourceError, match="^IMPORT_XML_LIMIT$"):
        read(values)
    monkeypatch.setattr(workbook, "MAX_PARTS", 4)
    with pytest.raises(ImportSourceError, match="^IMPORT_ZIP_LIMIT$"):
        read()


def test_combined_expansion_and_total_xml_nodes_are_bounded(monkeypatch):
    values = parts()
    size = sum(len(value.encode("utf-8")) for value in values.values())
    monkeypatch.setattr(workbook, "MAX_EXPANDED_BYTES", size)
    assert read(values).rows
    monkeypatch.setattr(workbook, "MAX_EXPANDED_BYTES", size - 1)
    with pytest.raises(ImportSourceError, match="^IMPORT_ZIP_LIMIT$"):
        read(values)
    monkeypatch.setattr(workbook, "MAX_EXPANDED_BYTES", size)
    monkeypatch.setattr(workbook, "MAX_XML_NODES", 10)
    with pytest.raises(ImportSourceError, match="^IMPORT_XML_LIMIT$"):
        read(values)


def test_malformed_cell_structure_is_rejected_instead_of_discarding_its_value():
    values = parts()
    values["xl/worksheets/sheet1.xml"] = values["xl/worksheets/sheet1.xml"].replace('r="B2" t="inlineStr"', 'r="B2" t="n"')
    with pytest.raises(ImportSourceError, match="^IMPORT_XLSX_CELL$"):
        read(values)


@pytest.mark.parametrize(("old", "new", "code"), [
    ('r="B2"', 'r="A2"', "IMPORT_XLSX_CELL"),
    ('r="B2"', 'r="B3"', "IMPORT_XLSX_CELL"),
    ('r="B2"', 'r="AG2"', "IMPORT_COLUMN_LIMIT"),
    ('r="B2"', 'r="D2"', "IMPORT_ROW_WIDTH"),
    ('<row r="2">', '<row r="999999">', "IMPORT_ROW_LIMIT"),
    ('<row r="2">', '<row r="1">', "IMPORT_XLSX_STRUCTURE"),
    ('</worksheet>', '<mergeCells><mergeCell ref="A1:B1"/></mergeCells></worksheet>', "IMPORT_XLSX_STRUCTURE"),
])
def test_rejects_ambiguous_cells_and_unbounded_coordinates(old, new, code):
    values = parts()
    values["xl/worksheets/sheet1.xml"] = values["xl/worksheets/sheet1.xml"].replace(old, new)
    with pytest.raises(ImportSourceError, match="^" + code + "$"):
        read(values)


def test_sheet_relationships_must_point_to_distinct_internal_worksheets():
    for target in ["../worksheets/sheet1.xml", "https://example.invalid/file.xml", "worksheets/missing.xml"]:
        values = parts()
        values["xl/_rels/workbook.xml.rels"] = values["xl/_rels/workbook.xml.rels"].replace("worksheets/sheet1.xml", target)
        with pytest.raises(ImportSourceError, match="^IMPORT_XLSX_STRUCTURE$"):
            read(values)
    values = parts()
    values["xl/_rels/workbook.xml.rels"] = values["xl/_rels/workbook.xml.rels"].replace("worksheets/sheet1.xml", "/xl/worksheets/sheet1.xml")
    assert read(values).rows[0].values[0] == "RWT_0007_WOOD"


def test_invalid_xml_zip_and_macro_content_types_are_static_errors():
    for value in (b"not zip", archive({"[Content_Types].xml": "<broken"})):
        with pytest.raises(ImportSourceError, match="^IMPORT_XLSX_INVALID$"):
            workbook.read_xlsx(value, sheet="Materials")
    values = parts()
    values["[Content_Types].xml"] = values["[Content_Types].xml"].replace("spreadsheetml.sheet.main", "ms-excel.sheet.macroEnabled.main")
    with pytest.raises(ImportSourceError, match="^IMPORT_ACTIVE_CONTENT$"):
        read(values)


def test_workbooks_created_by_an_independent_writer_have_compatible_literal_values():
    from openpyxl import Workbook
    from openpyxl.styles import Font
    output = io.BytesIO()
    document = Workbook()
    sheet = document.active
    sheet.title = "Historical PBR"
    sheet.append(["Identity", "Name", "Project", "Brand", "Processor", "Note"])
    sheet.append(["RWT_0007_G03", "Dub & kámen", "Explicit project", "Explicit brand", "Explicit processor", None])
    sheet.append(["RWT_0012_G04", "  Travertine  ", "Second project", "Explicit brand", "Explicit processor", 12.5])
    sheet["A1"].font = Font(bold=True)
    sheet.column_dimensions["B"].width = 30
    document.create_sheet("Unused")
    document.save(output)
    data = output.getvalue()
    assert workbook.xlsx_sheets(data) == ("Historical PBR", "Unused")
    table = workbook.read_xlsx(data, sheet="Historical PBR")
    assert table.rows[0].values == ("RWT_0007_G03", "Dub & kámen", "Explicit project", "Explicit brand", "Explicit processor", "")
    assert table.rows[1].values[1] == "  Travertine  "
    assert table.rows[1].values[-1] == "12.5"


@pytest.mark.parametrize("active", ["formula", "hyperlink"])
def test_independently_written_formula_and_external_link_are_rejected(active):
    from openpyxl import Workbook
    output = io.BytesIO()
    document = Workbook()
    sheet = document.active
    sheet.append(["Identity", "Name"])
    sheet.append(["RWT_0001_G03", "Literal"])
    if active == "formula":
        sheet["B2"] = "=1+1"
    else:
        sheet["B2"].hyperlink = "https://example.invalid/synthetic"
    document.save(output)
    with pytest.raises(ImportSourceError, match="^IMPORT_(FORMULA|EXTERNAL_LINK)$"):
        workbook.read_xlsx(output.getvalue(), sheet=sheet.title)
